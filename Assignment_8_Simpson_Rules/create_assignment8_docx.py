import os
from pathlib import Path
import math
import matplotlib.pyplot as plt
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def generate_comparison_plots(output_dir: Path):
    """Generate 3-method comparison plots for Question 1 and Question 2."""
    output_dir.mkdir(parents=True, exist_ok=True)
    
    # ------------------ Plot 1: Q1 Error Comparison ------------------
    trap_n = [1, 2, 4, 8, 16]
    trap_err = [1.3333333333, 0.3333333333, 0.0833333333, 0.0208333333, 0.0052083333]
    simp13_n = [2, 4, 8, 16]
    simp13_err = [1e-15, 1e-15, 1e-15, 1e-15]
    simp38_n = [3, 6, 12]
    simp38_err = [1e-15, 1e-15, 1e-15]

    plt.figure(figsize=(7, 4.2))
    plt.semilogy(trap_n, trap_err, 'o-', color='#1f77b4', linewidth=2, label='Trapezoidal Rule (O(h²))')
    plt.semilogy(simp13_n, simp13_err, 's--', color='#2ca02c', linewidth=2, label='Simpson 1/3 Rule (Exact: Error ~ 0)')
    plt.semilogy(simp38_n, simp38_err, '^-.', color='#d62728', linewidth=2, label='Simpson 3/8 Rule (Exact: Error ~ 0)')
    plt.axhline(1e-15, color='gray', linestyle=':', label='Machine Precision Floor (10⁻¹⁵)')
    plt.title('Question 1: Error Comparison of the 3 Methods (f(x) = x²)', fontsize=11, fontweight='bold', pad=10)
    plt.xlabel('Number of Intervals (n)', fontsize=10.5)
    plt.ylabel('Absolute Error (Log Scale)', fontsize=10.5)
    plt.grid(True, which="both", ls="--", alpha=0.5)
    plt.legend(fontsize=9, loc='upper right')
    plot1_path = output_dir / "three_methods_q1_error.png"
    plt.savefig(plot1_path, dpi=200, bbox_inches='tight')
    plt.close()

    # ------------------ Plot 2: Q2 Error vs Step Size (Log-Log) ------------------
    trap_h = [0.25, 0.125, 0.0625, 0.03125]
    trap_q2_err = [0.0038400350, 0.0009585180, 0.0002395360, 0.0000598782]
    s13_h = [0.25, 0.125, 0.0625, 0.03125]
    s13_q2_err = [0.0000312470, 0.0000019877, 0.0000001246, 0.0000000078]
    s38_h = [0.333333, 0.166667, 0.083333, 0.041667]
    s38_q2_err = [0.0001681868, 0.0000139247, 0.0000008838, 0.0000000554]

    plt.figure(figsize=(7, 4.2))
    plt.loglog(trap_h, trap_q2_err, 'o-', color='#1f77b4', linewidth=2, label='Trapezoidal Rule (Slope = 2, O(h²))')
    plt.loglog(s13_h, s13_q2_err, 's-', color='#2ca02c', linewidth=2, label='Simpson 1/3 Rule (Slope = 4, O(h⁴))')
    plt.loglog(s38_h, s38_q2_err, '^-', color='#d62728', linewidth=2, label='Simpson 3/8 Rule (Slope = 4, O(h⁴))')
    plt.title('Question 2: Log-Log Error vs Step Size h (f(x) = exp(-x²))', fontsize=11, fontweight='bold', pad=10)
    plt.xlabel('Step Size (h)', fontsize=10.5)
    plt.ylabel('Absolute Error', fontsize=10.5)
    plt.grid(True, which="both", ls="--", alpha=0.5)
    plt.legend(fontsize=9, loc='lower right')
    plot2_path = output_dir / "three_methods_q2_loglog.png"
    plt.savefig(plot2_path, dpi=200, bbox_inches='tight')
    plt.close()

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=110, right=110):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D5DD", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_assignment_8_doc():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header setup
        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = header_p.add_run("Numerical Computing II — Assignment VIII | Kunal Kamod (MS2513)")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(120, 120, 120)

    base_dir = Path(__file__).resolve().parent
    nested_dir = base_dir / "Assignment_8_Simpson_Rules"
    if nested_dir.exists():
        img_base = nested_dir / "output"
    else:
        img_base = base_dir / "output"

    generate_comparison_plots(img_base)

    # Base styling
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(35, 35, 35)

    # Title Block
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_t1 = p_title.add_run("Numerical Computing II\n")
    run_t1.bold = True
    run_t1.font.size = Pt(15)
    run_t1.font.color.rgb = RGBColor(20, 30, 55)

    run_t2 = p_title.add_run("Assignment VIII — Simpson's 1/3 and 3/8 Rules with Three-Method Comparison\n")
    run_t2.bold = True
    run_t2.font.size = Pt(12)
    run_t2.font.color.rgb = RGBColor(40, 60, 90)

    run_t3 = p_title.add_run("Date: Oct 11, 2026\n\n")
    run_t3.font.size = Pt(9.5)
    run_t3.font.color.rgb = RGBColor(100, 100, 100)

    run_name = p_title.add_run("Name: Kunal Kamod\nRoll No: MS2513")
    run_name.bold = True
    run_name.font.size = Pt(10.5)

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(6)
    p_intro.paragraph_format.space_after = Pt(6)
    p_intro.paragraph_format.line_spacing = 1.15
    p_intro.add_run(
        "For this assignment, I implemented Simpson's 1/3 Rule and Simpson's 3/8 Rule in Python to evaluate definite integrals, "
        "verify their fourth-order accuracy O(h⁴), and compare both methods with the Trapezoidal Rule."
    )

    def add_divider():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(8)
        pPr = p._p.get_or_add_pPr()
        pbdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="D0D5DD"/></w:pBdr>')
        pPr.append(pbdr)

    add_divider()

    # =========================================================================
    # PART 1: SIMPSON'S 1/3 RULE
    # =========================================================================
    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.space_before = Pt(2)
    p_p1.paragraph_format.space_after = Pt(2)
    r_p1 = p_p1.add_run("PART 1: SIMPSON'S 1/3 RULE")
    r_p1.bold = True
    r_p1.font.size = Pt(13)
    r_p1.font.color.rgb = RGBColor(20, 80, 150)

    p_m13 = doc.add_paragraph()
    p_m13.paragraph_format.space_before = Pt(2)
    p_m13.paragraph_format.space_after = Pt(6)
    p_m13.paragraph_format.line_spacing = 1.15
    p_m13.add_run("Method & Formula:\n").bold = True
    p_m13.add_run(
        "• Logic: Fits a parabola (quadratic curve) through sets of 3 points (each spanning 2 strips). Therefore, total intervals n must be even.\n"
        "• Formula:\n"
        "       I ≈ (h / 3) · [ f(x₀) + 4 ∑_{i odd} f(x_i) + 2 ∑_{i even} f(x_i) + f(x_n) ]\n"
        "• Why the weights (1, 4, 2, 4, ..., 1): Midpoints inside each parabolic curve get weight 4; boundaries shared between parabolas add to 2 (1 + 1)."
    )

    # Q1 Simpson 1/3
    p_q1_13 = doc.add_paragraph()
    p_q1_13.paragraph_format.space_before = Pt(4)
    p_q1_13.paragraph_format.space_after = Pt(3)
    p_q1_13.add_run("Question 1: I = ∫₀² x² dx (Exact = 8/3 ≈ 2.6666666667)").bold = True

    t1_13 = doc.add_table(rows=5, cols=5)
    t1_13.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1_13)
    h1 = ["n", "h", "Approximation", "Absolute Error", "Order"]
    for i, t in enumerate(h1):
        c = t1_13.rows[0].cells[i]
        c.text = t
        set_cell_background(c, "EAEFF8")
        set_cell_margins(c, 70, 70, 90, 90)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(20, 40, 80)

    q1_13_data = [
        ["2", "1.000000", "2.6666666667", "0.0000000000", "—"],
        ["4", "0.500000", "2.6666666667", "0.0000000000", "Exact"],
        ["8", "0.250000", "2.6666666667", "0.0000000000", "Exact"],
        ["16", "0.125000", "2.6666666667", "0.0000000000", "Exact"]
    ]
    for row_idx, vals in enumerate(q1_13_data, start=1):
        for col_idx, val in enumerate(vals):
            c = t1_13.rows[row_idx].cells[col_idx]
            c.text = val
            set_cell_background(c, "F9FAFB" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, 50, 50, 90, 90)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(9)

    # Q1 Graph 1/3
    q1_13_img = img_base / "simpson_1_3rd" / "question1_simpson_13_error.png"
    if q1_13_img.exists():
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(q1_13_img), width=Inches(5.0))

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(2)
        p_desc.paragraph_format.space_after = Pt(4)
        p_desc.paragraph_format.line_spacing = 1.15
        p_desc.add_run("What this graph tells us: ").bold = True
        p_desc.add_run(
            "The error is identically 0.0000000000 across all step sizes. Because Simpson's 1/3 rule fits parabolas and the integrand f(x) = x² is already a parabola, "
            "it captures the true area exactly with zero truncation error."
        )

    # Q2 Simpson 1/3
    p_q2_13 = doc.add_paragraph()
    p_q2_13.paragraph_format.space_before = Pt(4)
    p_q2_13.paragraph_format.space_after = Pt(3)
    p_q2_13.add_run("Question 2: I = ∫₀¹ e^(-x²) dx (Reference = 0.7468241328)").bold = True

    t2_13 = doc.add_table(rows=5, cols=5)
    t2_13.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2_13)
    h2 = ["n", "h", "Approximation", "Absolute Error", "Order"]
    for i, t in enumerate(h2):
        c = t2_13.rows[0].cells[i]
        c.text = t
        set_cell_background(c, "EAEFF8")
        set_cell_margins(c, 70, 70, 90, 90)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(20, 40, 80)

    q2_13_data = [
        ["4", "0.250000", "0.7468553798", "0.0000312470", "—"],
        ["8", "0.125000", "0.7468261205", "0.0000019877", "3.974534"],
        ["16", "0.062500", "0.7468242574", "0.0000001246", "3.995465"],
        ["32", "0.031250", "0.7468241406", "0.0000000078", "3.998963"]
    ]
    for row_idx, vals in enumerate(q2_13_data, start=1):
        for col_idx, val in enumerate(vals):
            c = t2_13.rows[row_idx].cells[col_idx]
            c.text = val
            set_cell_background(c, "F9FAFB" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, 50, 50, 90, 90)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(9)

    # Q2 Graph 1/3
    q2_13_img = img_base / "simpson_1_3rd" / "question2_simpson_13_error.png"
    if q2_13_img.exists():
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(q2_13_img), width=Inches(5.0))

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(2)
        p_desc.paragraph_format.space_after = Pt(4)
        p_desc.paragraph_format.line_spacing = 1.15
        p_desc.add_run("What this graph tells us: ").bold = True
        p_desc.add_run(
            "The error vs step size h drops in a steep straight line on a log-log scale with slope ≈ 4.0 (O(h⁴)). "
            "Halving h slashes the error by roughly 16 times, reaching an error of only 7.8 × 10⁻⁹ at n = 32."
        )

    # Q3 Simpson 1/3
    p_q3_13 = doc.add_paragraph()
    p_q3_13.paragraph_format.space_before = Pt(4)
    p_q3_13.paragraph_format.space_after = Pt(3)
    p_q3_13.paragraph_format.line_spacing = 1.15
    p_q3_13.add_run("Question 3: Experimental Data (n = 4 intervals, h = 0.5)\n").bold = True
    p_q3_13.add_run(
        "• Formula: I ≈ (h/3) · [ f(0) + 4·f(0.5) + 2·f(1.0) + 4·f(1.5) + f(2.0) ]\n"
        "• Calculation: I ≈ (0.5/3) · [ 1.00 + 4(1.65) + 2(2.70) + 4(4.48) + 7.39 ] = (0.5/3) · 38.31 = 6.3850\n"
        "• Estimated Integral = 6.3850000000 (Applicable because n = 4 is an even number of intervals)."
    )

    # Q3 Graph 1/3
    q3_13_img = img_base / "simpson_1_3rd" / "question3_data_graph.png"
    if q3_13_img.exists():
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(q3_13_img), width=Inches(5.0))

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(2)
        p_desc.paragraph_format.space_after = Pt(4)
        p_desc.paragraph_format.line_spacing = 1.15
        p_desc.add_run("What this graph tells us: ").bold = True
        p_desc.add_run(
            "The plot shows the experimental points. Simpson's 1/3 rule fits smooth parabolas across (0, 0.5, 1.0) and (1.0, 1.5, 2.0), "
            "providing a smoother and more accurate area estimate (6.3850) than straight-line interpolation."
        )

    add_divider()

    # =========================================================================
    # PART 2: SIMPSON'S 3/8 RULE
    # =========================================================================
    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.space_before = Pt(2)
    p_p2.paragraph_format.space_after = Pt(2)
    r_p2 = p_p2.add_run("PART 2: SIMPSON'S 3/8 RULE")
    r_p2.bold = True
    r_p2.font.size = Pt(13)
    r_p2.font.color.rgb = RGBColor(160, 40, 40)

    p_m38 = doc.add_paragraph()
    p_m38.paragraph_format.space_before = Pt(2)
    p_m38.paragraph_format.space_after = Pt(6)
    p_m38.paragraph_format.line_spacing = 1.15
    p_m38.add_run("Method & Formula:\n").bold = True
    p_m38.add_run(
        "• Logic: Fits a cubic (degree-3) polynomial across sets of 4 points (spanning 3 strips). Therefore, total intervals n must be a multiple of 3.\n"
        "• Formula:\n"
        "       I ≈ (3h / 8) · [ f(x₀) + 3·f(x₁) + 3·f(x₂) + 2·f(x₃) + 3·f(x₄) + 3·f(x₅) + ... + f(x_n) ]\n"
        "• Why the weights: Intermediate interior points get weight 3; boundaries shared between adjacent cubic segments add up to 2 (1 + 1)."
    )

    # Q1 Simpson 3/8
    p_q1_38 = doc.add_paragraph()
    p_q1_38.paragraph_format.space_before = Pt(4)
    p_q1_38.paragraph_format.space_after = Pt(3)
    p_q1_38.add_run("Question 1: I = ∫₀² x² dx (Exact = 8/3 ≈ 2.6666666667)").bold = True

    t1_38 = doc.add_table(rows=4, cols=5)
    t1_38.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1_38)
    for i, t in enumerate(h1):
        c = t1_38.rows[0].cells[i]
        c.text = t
        set_cell_background(c, "EAEFF8")
        set_cell_margins(c, 70, 70, 90, 90)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(20, 40, 80)

    q1_38_data = [
        ["3", "0.666667", "2.6666666667", "0.0000000000", "—"],
        ["6", "0.333333", "2.6666666667", "0.0000000000", "Exact"],
        ["12", "0.166667", "2.6666666667", "0.0000000000", "Exact"]
    ]
    for row_idx, vals in enumerate(q1_38_data, start=1):
        for col_idx, val in enumerate(vals):
            c = t1_38.rows[row_idx].cells[col_idx]
            c.text = val
            set_cell_background(c, "F9FAFB" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, 50, 50, 90, 90)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(9)

    # Q1 Graph 3/8
    q1_38_img = img_base / "simpson_3_8th" / "question1_simpson_38_error.png"
    if q1_38_img.exists():
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(q1_38_img), width=Inches(5.0))

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(2)
        p_desc.paragraph_format.space_after = Pt(4)
        p_desc.paragraph_format.line_spacing = 1.15
        p_desc.add_run("What this graph tells us: ").bold = True
        p_desc.add_run(
            "The absolute error is zero (within floating-point precision). Because Simpson's 3/8 rule uses cubic interpolation, "
            "it exactly integrates all polynomials of degree ≤ 3, including f(x) = x²."
        )

    # Q2 Simpson 3/8
    p_q2_38 = doc.add_paragraph()
    p_q2_38.paragraph_format.space_before = Pt(4)
    p_q2_38.paragraph_format.space_after = Pt(3)
    p_q2_38.add_run("Question 2: I = ∫₀¹ e^(-x²) dx (Reference = 0.7468241328)").bold = True

    t2_38 = doc.add_table(rows=5, cols=5)
    t2_38.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2_38)
    for i, t in enumerate(h2):
        c = t2_38.rows[0].cells[i]
        c.text = t
        set_cell_background(c, "EAEFF8")
        set_cell_margins(c, 70, 70, 90, 90)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(20, 40, 80)

    q2_38_data = [
        ["3", "0.333333", "0.7469923196", "0.0001681868", "—"],
        ["6", "0.166667", "0.7468380575", "0.0000139247", "3.594346"],
        ["12", "0.083333", "0.7468250167", "0.0000008838", "3.977713"],
        ["24", "0.041667", "0.7468241882", "0.0000000554", "3.995987"]
    ]
    for row_idx, vals in enumerate(q2_38_data, start=1):
        for col_idx, val in enumerate(vals):
            c = t2_38.rows[row_idx].cells[col_idx]
            c.text = val
            set_cell_background(c, "F9FAFB" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, 50, 50, 90, 90)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(9)

    # Q2 Graph 3/8
    q2_38_img = img_base / "simpson_3_8th" / "question2_simpson_38_error.png"
    if q2_38_img.exists():
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(q2_38_img), width=Inches(5.0))

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(2)
        p_desc.paragraph_format.space_after = Pt(4)
        p_desc.paragraph_format.line_spacing = 1.15
        p_desc.add_run("What this graph tells us: ").bold = True
        p_desc.add_run(
            "The log-log error curve slopes steeply at rate 4 (O(h⁴)), demonstrating that increasing n from 3 to 24 drops the error "
            "dramatically from 1.68 × 10⁻⁴ down to 5.54 × 10⁻⁸."
        )

    # Q3 Simpson 3/8 restriction
    p_q3_38 = doc.add_paragraph()
    p_q3_38.paragraph_format.space_before = Pt(4)
    p_q3_38.paragraph_format.space_after = Pt(3)
    p_q3_38.paragraph_format.line_spacing = 1.15
    p_q3_38.add_run("Question 3 Restriction for Simpson 3/8:\n").bold = True
    p_q3_38.add_run(
        "• The given dataset contains 5 points, giving n = 4 intervals.\n"
        "• Simpson 3/8 requires n to be a multiple of 3 (n = 3, 6, 9, ...).\n"
        "• Because 4 is not divisible by 3, pure Simpson 3/8 cannot be applied alone across the full interval [0, 2]."
    )

    # Q3 Graph 3/8
    q3_38_img = img_base / "simpson_3_8th" / "question3_data_graph.png"
    if q3_38_img.exists():
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(q3_38_img), width=Inches(5.0))

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(2)
        p_desc.paragraph_format.space_after = Pt(4)
        p_desc.paragraph_format.line_spacing = 1.15
        p_desc.add_run("What this graph tells us: ").bold = True
        p_desc.add_run(
            "This plot illustrates the 4 discrete intervals in the dataset. Because 4 cannot be partitioned into triplets of 3, "
            "pure Simpson 3/8 cannot cover all four segments without leaving an unintegrated remainder."
        )

    add_divider()

    # =========================================================================
    # PART 3: COMPARISON OF THE THREE METHODS
    # =========================================================================
    p_p3 = doc.add_paragraph()
    p_p3.paragraph_format.space_before = Pt(2)
    p_p3.paragraph_format.space_after = Pt(2)
    r_p3 = p_p3.add_run("PART 3: COMPARISON OF THE THREE METHODS (TRAPEZOIDAL vs 1/3 vs 3/8)")
    r_p3.bold = True
    r_p3.font.size = Pt(13)
    r_p3.font.color.rgb = RGBColor(20, 100, 40)

    p_comp_intro = doc.add_paragraph()
    p_comp_intro.paragraph_format.space_before = Pt(2)
    p_comp_intro.paragraph_format.space_after = Pt(6)
    p_comp_intro.paragraph_format.line_spacing = 1.15
    p_comp_intro.add_run(
        "Below is the complete side-by-side comparison of the Trapezoidal Rule, Simpson's 1/3 Rule, and Simpson's 3/8 Rule across all three assignment problems:"
    )

    # Comparison Table
    t_comp = doc.add_table(rows=7, cols=4)
    t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_comp)

    comp_headers = ["Feature / Metric", "Trapezoidal Rule", "Simpson's 1/3 Rule", "Simpson's 3/8 Rule"]
    for i, t in enumerate(comp_headers):
        c = t_comp.rows[0].cells[i]
        c.text = t
        set_cell_background(c, "EAEFF8")
        set_cell_margins(c, 70, 70, 90, 90)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(20, 40, 80)

    comp_rows = [
        ["Approximating Curve", "Linear (Degree 1 chords)", "Parabolic (Degree 2)", "Cubic (Degree 3)"],
        ["Theoretical Order", "O(h²) [Second-order]", "O(h⁴) [Fourth-order]", "O(h⁴) [Fourth-order]"],
        ["Interval Constraint", "Any n ≥ 1", "n must be even (2, 4, 6...)", "n multiple of 3 (3, 6, 9...)"],
        ["Q1 Error (f(x) = x²)", "0.0833 (at n=4)", "0.0000 (Exact for all n)", "0.0000 (Exact for all n)"],
        ["Q2 Error (at h ≈ 0.125)", "9.58 × 10⁻⁴ (at n=8)", "1.99 × 10⁻⁶ (at n=8)", "1.39 × 10⁻⁵ (at n=6, h=0.16)"],
        ["Q3 Result (n = 4)", "6.5125 (Applicable)", "6.3850 (Applicable, n even)", "Not directly applicable (n≠3k)"]
    ]

    for row_idx, vals in enumerate(comp_rows, start=1):
        for col_idx, val in enumerate(vals):
            c = t_comp.rows[row_idx].cells[col_idx]
            c.text = val
            set_cell_background(c, "F9FAFB" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, 50, 50, 90, 90)
            p = c.paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.runs[0].bold = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(8.5)

    # Comparison Graph 1A: Question 1 Approximation Comparison
    q1_prog_img = img_base / "question1_comparison.png"
    if q1_prog_img.exists():
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(q1_prog_img), width=Inches(5.0))

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(2)
        p_desc.paragraph_format.space_after = Pt(4)
        p_desc.paragraph_format.line_spacing = 1.15
        p_desc.add_run("What Graph 1 (Q1 Approximations) tells us: ").bold = True
        p_desc.add_run(
            "This plot compares the calculated integral values against the exact 2.6667 value. "
            "Both Simpson 1/3 and 3/8 immediately match the exact line across all n, whereas Trapezoidal starts with overestimation and slowly converges."
        )

    # Comparison Graph 1B: Question 1 Error Comparison (Trapezoid vs 1/3 vs 3/8)
    q1_comp_img = img_base / "three_methods_q1_error.png"
    if q1_comp_img.exists():
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(q1_comp_img), width=Inches(5.0))

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(2)
        p_desc.paragraph_format.space_after = Pt(4)
        p_desc.paragraph_format.line_spacing = 1.15
        p_desc.add_run("What Graph 2 (Q1 Three-Method Error) tells us: ").bold = True
        p_desc.add_run(
            "On a logarithmic error scale, the Trapezoidal Rule displays positive non-zero errors descending with slope 2, "
            "while both Simpson rules sit flat at machine precision 10⁻¹⁵ (0.0000000000 error)."
        )

    # Comparison Graph 2A: Question 2 Approximation Comparison
    q2_prog_img = img_base / "question2_comparison.png"
    if q2_prog_img.exists():
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(q2_prog_img), width=Inches(5.0))

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(2)
        p_desc.paragraph_format.space_after = Pt(4)
        p_desc.paragraph_format.line_spacing = 1.15
        p_desc.add_run("What Graph 3 (Q2 Approximations) tells us: ").bold = True
        p_desc.add_run(
            "This plot shows numerical convergence toward the reference value 0.746824. "
            "Both Simpson curves latch onto the reference horizontal line much faster than the Trapezoidal curve."
        )

    # Comparison Graph 2B: Question 2 Log-Log Error Comparison (Three Methods)
    q2_comp_img = img_base / "three_methods_q2_loglog.png"
    if q2_comp_img.exists():
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(q2_comp_img), width=Inches(5.0))

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(2)
        p_desc.paragraph_format.space_after = Pt(4)
        p_desc.paragraph_format.line_spacing = 1.15
        p_desc.add_run("What Graph 4 (Q2 Three-Method Log-Log Error) tells us: ").bold = True
        p_desc.add_run(
            "The slopes reveal the order of convergence: Trapezoidal has slope 2 (O(h²)), while Simpson 1/3 and 3/8 have slope 4 (O(h⁴)). "
            "At h = 0.03125, Simpson 1/3 attains an error of 7.8 × 10⁻⁹, whereas Trapezoidal error is still 6.0 × 10⁻⁵—roughly 7,700 times higher."
        )

    # Final Summary
    p_summary = doc.add_paragraph()
    p_summary.paragraph_format.space_before = Pt(4)
    p_summary.paragraph_format.space_after = Pt(6)
    p_summary.paragraph_format.line_spacing = 1.15
    p_summary.add_run("Logical Conclusions & When to Use Which Method:\n").bold = True
    p_summary.add_run(
        "1. Trapezoidal Rule is the most flexible: It works for any number of intervals n without restriction, but needs a very fine step size h due to O(h²) accuracy.\n"
        "2. Simpson's 1/3 Rule is the most efficient choice: It achieves fourth-order O(h⁴) accuracy with simple alternating weights (1, 4, 2, ...), provided n is even.\n"
        "3. Simpson's 3/8 Rule is equally accurate (O(h⁴)): It serves as an essential companion when n is a multiple of 3, or when combined with 1/3 to integrate odd numbers of intervals."
    )

    doc_path = base_dir / "Assignment_8_Simpson_Rules_Writeup.docx"
    doc.save(str(doc_path))
    print(f"Document successfully created at: {doc_path}")

    if nested_dir.exists():
        doc.save(str(nested_dir / "Assignment_8_Simpson_Rules_Writeup.docx"))
        print(f"Also saved in: {nested_dir / 'Assignment_8_Simpson_Rules_Writeup.docx'}")

if __name__ == "__main__":
    build_assignment_8_doc()
